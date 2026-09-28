import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

from job_signal_viewer import Explorer


class WorkflowTests(unittest.TestCase):
    def test_import_update_error_and_export(self):
        app = Explorer()
        app.withdraw()

        def wait_for_worker():
            deadline = time.monotonic() + 15
            while app.busy and time.monotonic() < deadline:
                app.update()
                time.sleep(0.02)
            self.assertFalse(app.busy, 'worker timeout')

        try:
            with tempfile.TemporaryDirectory() as folder:
                root = Path(folder).resolve()
                first = root / 'a.job'
                second = root / 'b.JOB'
                first.write_text('do40=1 #전진\nwait di40 #확인', encoding='utf-8')
                second.write_bytes('do40=0 #정지'.encode('cp949'))
                app.load_paths([str(first), str(second)])
                wait_for_worker()
                self.assertEqual((len(app.documents), len(app.signals)), (2, 3))
                self.assertEqual(str(app.export_button['state']), 'normal')

                first.write_text('do41=1 #갱신', encoding='utf-8')
                app.load_paths([str(first)])
                wait_for_worker()
                self.assertEqual((len(app.documents), len(app.signals)), (2, 2))
                self.assertIn('DO41', app.by_signal)

                with patch('job_signal_viewer.messagebox.showerror') as error:
                    app.load_paths([str(root / 'missing.job')])
                    wait_for_worker()
                    self.assertTrue(error.called)
                    self.assertEqual(len(app.documents), 2)

                output = root / 'saved.xlsx'
                app.search.set('아무것도 없음')
                with patch('job_signal_viewer.filedialog.asksaveasfilename', return_value=str(output)), \
                     patch('job_signal_viewer.messagebox.showinfo') as info:
                    app.save_excel()
                    wait_for_worker()
                    self.assertTrue(info.called)
                wb = load_workbook(output)
                try:
                    self.assertEqual(wb['사용 내역'].max_row, 3)
                    for row in range(2, 4):
                        self.assertTrue(wb['사용 내역'][f'F{row}'].hyperlink.location.startswith("'원본 텍스트'!"))
                finally:
                    wb.close()
        finally:
            for callback in app.tk.call('after', 'info'):
                app.after_cancel(callback)
            app.destroy()


if __name__ == '__main__':
    unittest.main()
