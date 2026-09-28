import sys
from pathlib import Path
import tempfile
import unittest
ROOT = Path(__file__).resolve().parents[1]
SAMPLE = Path(__file__).resolve().parent / 'fixtures' / 'sample.job'
sys.path.insert(0, str(ROOT / 'src'))
from job_signal_extractor import parse_text, read_job, collect_files, extract
from openpyxl import load_workbook

class ExtractorTests(unittest.TestCase):
    def test_tokens(self):
        rows = parse_text('do40 = 0 #전진 OFF\nwait di40 #확인\nif DI73==1 and di74!=0 #조건\nwait do31')
        self.assertEqual([r.name for r in rows], ['DO40','DI40','DI73','DI74','DO31'])
        self.assertEqual([r.value for r in rows], [0,None,1,0,None])
        self.assertEqual(rows[2].operator, '==')
        self.assertEqual(rows[3].comment, '조건')

    def test_comments_strings_boundaries(self):
        rows = parse_text('# di99\nprint "di3 # do4"\nmydi2 di4_extra do5x\nwait di7 #이름 # di88\nprint "hi"; do8=1')
        self.assertEqual([r.name for r in rows], ['DI7', 'DO8'])
        self.assertEqual(rows[0].comment, '이름 # di88')
        self.assertEqual(rows[0].line, 4)
        self.assertIsNone(parse_text('if di2==10')[0].value)
        self.assertIsNone(parse_text('do2=1+v')[0].value)

    def test_files_and_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            a = root / 'a.job'
            a.write_bytes('do1=0 #한글\nwait di2\ndo1=1 #=SUM(A1)'.encode('cp949'))
            b = root / 'b.JOB'
            b.write_text('# empty', encoding='utf-8')
            self.assertEqual(read_job(a)[1], 'cp949')
            self.assertEqual(len(collect_files([str(root), str(a)])), 2)
            out = root / 'out.xlsx'
            self.assertEqual(extract([str(root)], out), (2,3,2))
            wb = load_workbook(out)
            self.assertEqual(wb['사용 내역']['F4'].value, '=SUM(A1)')
            self.assertEqual(wb['사용 내역']['F4'].data_type, 's')
            self.assertEqual(wb['처리 파일']['D3'].value, 0)
            self.assertEqual(wb['신호 요약']['D3'].value, '한글\n=SUM(A1)')
            self.assertEqual(wb['사용 내역'].auto_filter.ref, 'A1:K4')
            wb.close()

    def test_sample(self):
        text, enc = read_job(SAMPLE)
        rows = parse_text(text)
        self.assertEqual((len(rows), len({r.name for r in rows})), (18,18))
        self.assertEqual(rows[0].name, 'DO31')
        self.assertEqual(rows[0].comment, '로봇 원위치 확인')
        self.assertEqual(enc, 'utf-8-sig')
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'sample.xlsx'
            self.assertEqual(extract([str(SAMPLE)], output), (1, 18, 18))
            wb = load_workbook(output)
            try:
                self.assertEqual(wb['사용 내역'].max_row, 19)
                self.assertEqual(wb['신호 요약'].max_row, 19)
                self.assertEqual(sum(row[4] for row in wb['신호 요약'].iter_rows(min_row=2, values_only=True)), 18)
            finally:
                wb.close()

if __name__ == '__main__':
    unittest.main()
