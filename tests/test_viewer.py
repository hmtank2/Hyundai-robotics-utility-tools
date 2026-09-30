import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = Path(__file__).resolve().parent / 'fixtures' / 'sample.job'
sys.path.insert(0, str(ROOT / 'src'))

from job_signal_extractor import export_excel, parse_text, read_job
from job_signal_viewer import Explorer


class ViewerTests(unittest.TestCase):
    def test_job_call_graph_and_source_jump(self):
        folder = Path(__file__).resolve().parent / 'fixtures' / 'calls'
        app = Explorer()
        app.withdraw()
        try:
            for path in sorted(folder.glob('*.job')):
                text, encoding = read_job(path)
                key = str(path.resolve())
                app.documents[key] = (text, encoding, parse_text(text, key))
            app.reindex()
            app.update()
            main = str((folder / 'main.job').resolve())
            target = str((folder / '9250.job').resolve())
            app.show_node(('file', main))
            self.assertEqual(len(app.by_file[main]), 0)
            self.assertEqual(len(app.calls_out[main]), 3)
            self.assertEqual(len(app.neighbors()), 3)
            self.assertEqual({neighbor.relation for neighbor in app.neighbors()}, {'out', 'missing'})
            self.assertEqual(len(app.call_details.get_children()), 3)
            app.call_details.selection_set('0')
            app.select_call()
            self.assertIn('2행', app.source_title.get())
            self.assertTrue(app.source.tag_ranges('target'))
            app.show_node(('file', target))
            self.assertEqual(len(app.calls_in[target]), 2)
            self.assertEqual(len(app.neighbors()), 3)  # DI40 plus two calling JOBs
        finally:
            for callback in app.tk.call('after', 'info'):
                app.after_cancel(callback)
            app.destroy()

    def test_navigation_and_export(self):
        app = Explorer()
        app.withdraw()
        try:
            text, encoding = read_job(SAMPLE)
            path = str(SAMPLE.resolve())
            other = str((SAMPLE.parent / 'other' / SAMPLE.name).resolve())
            extra = 'do40=1 #두 번째 프로그램\nwait di40 #공통 입력'
            app.documents = {
                path: (text, encoding, parse_text(text, path)),
                other: (extra, 'utf-8', parse_text(extra, other)),
            }
            app.reindex()
            app.update()
            app.show_node(('signal', 'DO40'))
            self.assertEqual(len(app.neighbors()), 2)
            app.details.selection_set(str(next(i for i, row in enumerate(app.current_rows) if row.file == other)))
            app.select_occurrence()
            self.assertIn('두 번째 프로그램', app.source.get('1.0', 'end'))
            self.assertTrue(app.source.tag_ranges('target'))
            app.search.set('공통 입력')
            self.assertEqual(list(app.signal_tree.get_children()), ['DI40'])
            app.show_node(('file', path))
            self.assertEqual(len(app.neighbors()), 18)
            app.page(1)
            self.assertEqual(app.graph_page, 1)
            app.search.set('')

            with tempfile.TemporaryDirectory() as folder:
                output = Path(folder) / 'multifile.xlsx'
                export_excel(
                    app.signals,
                    [(Path(p).name, p, data[1], len(data[2])) for p, data in app.documents.items()],
                    output,
                    {p: data[0] for p, data in app.documents.items()},
                )
                wb = load_workbook(output)
                try:
                    sample_row = next(i for i, row in enumerate(app.signals, 2) if row.file == path)
                    self.assertEqual(wb['사용 내역'][f'F{sample_row}'].hyperlink.location, "'원본 텍스트'!C4")
                    self.assertEqual(wb['원본 텍스트']['C4'].value.strip(), 'wait do31 #로봇 원위치 확인')
                    self.assertEqual(sum(row[4] for row in wb['신호 요약'].iter_rows(min_row=2, values_only=True)), 20)
                finally:
                    wb.close()
            app.clear()
            self.assertFalse(app.documents)
            self.assertFalse(app.signal_tree.get_children())
        finally:
            for callback in app.tk.call('after', 'info'):
                app.after_cancel(callback)
            app.destroy()


if __name__ == '__main__':
    unittest.main()
