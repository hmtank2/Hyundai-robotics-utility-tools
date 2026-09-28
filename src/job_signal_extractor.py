"""현대 로봇 JOB의 DI/DO와 같은 줄의 # 주석을 XLSX로 추출. Python 3.10+."""
from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import re
import sys
import tempfile
import os

# 다른 변수 이름 내부의 di/do는 제외. 대소문자는 구분하지 않습니다.
SIGNAL_RE = re.compile(r"(?<![\w.])(?P<kind>di|do)(?P<number>[0-9]+)(?!\w)", re.I)
VALUE_RE = re.compile(r"\s*(==|!=|<>|<=|>=|=|<|>)\s*([01])(?![\w.]|\s*[+*/-])")


@dataclass
class Signal:
    kind: str
    number: int
    operator: str
    value: int | None
    comment: str
    file: str
    line: int
    code: str
    raw: str

    @property
    def name(self):
        return f"{self.kind}{self.number}"


def split_line(line: str) -> tuple[str, str]:
    """따옴표 안의 문자열은 검색에서 제외하고, 밖의 첫 #에서 주석 분리."""
    masked = list(line)
    quote = None
    i = 0
    while i < len(line):
        char = line[i]
        if quote:
            masked[i] = " "
            if char == "\\" and i + 1 < len(line):
                i += 1
                masked[i] = " "
            elif char == quote:
                if i + 1 < len(line) and line[i + 1] == quote:
                    i += 1
                    masked[i] = " "
                else:
                    quote = None
        elif char in ('"', "'"):
            quote = char
            masked[i] = " "
        elif char == "#":
            return "".join(masked[:i]), line[i + 1:].strip()
        i += 1
    return "".join(masked), ""


def parse_text(text: str, source: str = "") -> list[Signal]:
    results = []
    for line_no, raw in enumerate(text.splitlines(), 1):
        code, comment = split_line(raw)
        for match in SIGNAL_RE.finditer(code):
            value = VALUE_RE.match(code, match.end())
            results.append(Signal(
                match['kind'].upper(), int(match['number']),
                value[1] if value else "", int(value[2]) if value else None,
                comment, source, line_no, raw[:len(code)].strip(), raw,
            ))
    return results


def read_job(path: Path, encoding: str | None = None) -> tuple[str, str]:
    raw = path.read_bytes()
    if encoding:
        return raw.decode(encoding), encoding
    encodings = ["utf-8-sig", "cp949"]
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        encodings = ["utf-16"]
    for candidate in encodings:
        try:
            return raw.decode(candidate), candidate
        except UnicodeDecodeError:
            pass
    raise ValueError(f"인코딩을 판별할 수 없습니다: {path}\n--encoding 옵션을 지정하세요.")


def collect_files(inputs: list[str], recursive: bool = False) -> list[Path]:
    files = set()
    for item in inputs:
        path = Path(item).expanduser().resolve()
        if path.is_dir():
            candidates = path.rglob('*') if recursive else path.iterdir()
            files.update(p.resolve() for p in candidates if p.is_file() and p.suffix.lower() == '.job')
        elif path.is_file() and path.suffix.lower() == '.job':
            files.add(path)
        else:
            raise ValueError(f"JOB 파일 또는 폴더가 아닙니다: {path}")
    if not files:
        raise ValueError("선택한 경로에 JOB 파일이 없습니다.")
    return sorted(files, key=lambda p: str(p).casefold())


def export_excel(signals: list[Signal], files: list[tuple], output: Path, sources: dict[str, str] | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            if any(isinstance(v, str) and len(v) > 32767 for v in row):
                raise ValueError('엑셀 셀 글자 수 한도를 넘었습니다. 파일을 나눠 처리하세요.')
            if ws.max_row >= 1048576:
                raise ValueError(f"{name}: 엑셀 행 수 한도를 넘었습니다. 파일을 나눠 처리하세요.")
            ws.append(row)
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value, str):
                    if len(cell.value) > 32767:
                        raise ValueError("엑셀 셀 글자 수 한도를 넘었습니다. 파일을 나눠 처리하세요.")
                    # JOB 주석이 '='로 시작해도 엑셀 수식으로 실행되지 않도록 저장.
                    cell.data_type = 's'
                cell.font = Font(name='맑은 고딕', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row > 1 and cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EFF5FA')
            height = max(str(c.value or '').count('\n') + 1 for c in cells)
            ws.row_dimensions[cells[0].row].height = min(409, max(32, 17 * height + 10))
        for c in ws[1]:
            c.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='183F61')
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = 'D2' if name == '신호 요약' else 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        return ws

    groups = defaultdict(list)
    for signal in signals:
        groups[(signal.kind, signal.number)].append(signal)
    summary = []
    for (kind, number), items in sorted(groups.items()):
        comments = list(dict.fromkeys(s.comment for s in items if s.comment))
        summary.append([kind, number, f'{kind}{number}', '\n'.join(comments),
                        len(items), len({s.file for s in items}), len(comments),
                        sum(not s.comment for s in items), '\n'.join(dict.fromkeys(s.file for s in items))])
    sheet('신호 요약', ['구분', '번호', '신호', '주석 목록 (원문)', '출현 횟수', '파일 수', '주석 종류 수', '주석 없는 횟수', '사용 파일 (전체 경로)'],
          summary, [9, 9, 12, 65, 14, 11, 15, 18, 85])
    detail = sheet('사용 내역', ['구분', '번호', '신호', '연산자', '명시 값', '주석', '파일명', '줄 번호', '명령문', '전체 경로', '원본 줄'],
          [[s.kind, s.number, s.name, s.operator, s.value, s.comment, Path(s.file).name,
            s.line, s.code, s.file, s.raw] for s in signals],
          [9, 9, 12, 10, 11, 55, 20, 11, 46, 65, 80])
    sheet('처리 파일', ['파일명', '전체 경로', '인코딩', '추출 건수'], files, [24, 95, 18, 14])
    if sources is not None:
        from openpyxl.worksheet.hyperlink import Hyperlink
        locations = {}
        source_rows = []
        for path, text in sources.items():
            for line_no, line in enumerate(text.splitlines(), 1):
                locations[(path, line_no)] = len(source_rows) + 2
                source_rows.append([Path(path).name, line_no, line, path])
        sheet('원본 텍스트', ['파일명', '줄 번호', '원본 텍스트 (불러온 시점)', '전체 경로'], source_rows, [24, 12, 120, 85])
        for row, signal in enumerate(signals, 2):
            target = locations.get((signal.file, signal.line))
            if target:
                for col in ('F', 'H'):
                    cell = detail[f'{col}{row}']
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'원본 텍스트'!C{target}")
                    cell.font = Font(name='맑은 고딕', size=11, color='1765B2', underline='single')
    output = output.resolve()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('출력 파일 확장자는 .xlsx여야 합니다.')
    output.parent.mkdir(parents=True, exist_ok=True)
    # 저장에 실패해도 기존 결과 파일을 손상시키지 않도록 임시 파일 사용.
    fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=output.parent)
    os.close(fd)
    try:
        wb.save(temp)
        os.replace(temp, output)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def extract(inputs, output, recursive=False, encoding=None):
    signals, files = [], []
    for path in collect_files(inputs, recursive):
        text, used_encoding = read_job(path, encoding)
        found = parse_text(text, str(path))
        signals.extend(found)
        files.append((path.name, str(path), used_encoding, len(found)))
    export_excel(signals, files, Path(output))
    return len(files), len(signals), len({s.name for s in signals})


def gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        paths = filedialog.askopenfilenames(title='추출할 JOB 파일 선택 (여러 개 선택 가능)',
                                           filetypes=[('로봇 JOB 파일', '*.job *.JOB'), ('모든 파일', '*.*')])
        if not paths:
            return
        output = filedialog.asksaveasfilename(title='엑셀 저장 위치', defaultextension='.xlsx',
                                            initialfile='JOB_신호목록.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not output:
            return
        count, occurrences, unique = extract(paths, output)
        messagebox.showinfo('추출 완료', f'{count}개 파일 / {unique}개 신호 / {occurrences}건\n\n{output}')
    except Exception as exc:
        messagebox.showerror('추출 실패', f'{exc}\n\nopenpyxl 설치 여부와 엑셀 파일이 열려 있는지 확인하세요.')
    finally:
        root.destroy()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('inputs', nargs='*', help='JOB 파일 또는 폴더 (여러 경로 가능)')
    parser.add_argument('-o', '--output', default='JOB_신호목록.xlsx')
    parser.add_argument('-r', '--recursive', action='store_true', help='하위 폴더도 검색')
    parser.add_argument('--encoding', help='인코딩 강제 지정, 예: cp949')
    parser.add_argument('--overwrite', action='store_true', help='기존 결과 덮어쓰기')
    args = parser.parse_args()
    if not args.inputs:
        gui()
        return 0
    try:
        if Path(args.output).exists() and not args.overwrite:
            raise ValueError('결과 파일이 이미 있습니다. 다른 이름 또는 --overwrite를 지정하세요.')
        files, occurrences, unique = extract(args.inputs, args.output, args.recursive, args.encoding)
        print(f'완료: 파일 {files}개, 신호 {unique}개, 출현 {occurrences}건 → {Path(args.output).resolve()}')
        return 0
    except Exception as exc:
        print(f'실패: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
